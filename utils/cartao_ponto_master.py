import csv
import json
import os
import re
import sys
from pathlib import Path

import cv2
import fitz
import numpy as np
from PIL import Image

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from utils.gemini.cartao_ponto import SYSTEM_PROMPT_CARTAO_PONTO, limpar_json
from utils.gemini_service import modelos_disponiveis
from utils.tesseract_config import pytesseract


EXT_LOGICO = {".csv", ".txt", ".xlsx", ".xlsm"}
EXT_IMAGEM = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
EXT_PDF = {".pdf"}

TIPOS_BATIDA = {
    "entrada": ["entrada", "ent", "entrada 1", "ent1", "1 entrada", "1a entrada", "1ª entrada"],
    "saida_intervalo": ["saida intervalo", "saída intervalo", "sai1", "saida 1", "saída 1", "1 saida", "1ª saída"],
    "retorno_intervalo": ["retorno", "retorno intervalo", "ent2", "entrada 2", "2 entrada", "2ª entrada"],
    "saida": ["saida", "saída", "sai2", "saida 2", "saída 2", "2 saida", "2ª saída"],
}


def normalizar_texto(valor):
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def normalizar_data(valor):
    texto = normalizar_texto(valor)
    match = re.search(r"\b([0-3]?\d)[/.\-]([01]?\d)[/.\-](\d{2,4})\b", texto)
    if not match:
        return ""

    dia, mes, ano = match.groups()
    if len(ano) == 2:
        ano = f"20{ano}"

    return f"{dia.zfill(2)}/{mes.zfill(2)}/{ano}"


def normalizar_horario(valor):
    texto = normalizar_texto(valor)
    match = re.search(r"\b([0-2]?\d)\s*[:hH;.]\s*([0-5]\d)\b", texto)
    if not match:
        return ""

    hora, minuto = match.groups()
    hora_int = int(hora)
    if hora_int > 23:
        return ""

    return f"{str(hora_int).zfill(2)}:{minuto}"


def detectar_tipo_coluna(nome_coluna):
    nome = normalizar_texto(nome_coluna).lower()
    nome = nome.replace("í", "i").replace("í", "i").replace("ã", "a").replace("ç", "c")

    if any(chave in nome for chave in ["data", "dia"]):
        return "data"

    for tipo, aliases in TIPOS_BATIDA.items():
        for alias in aliases:
            alias_norm = alias.replace("í", "i").replace("ã", "a").replace("ç", "c")
            if alias_norm in nome:
                return tipo

    if any(chave in nome for chave in ["nome", "funcionario", "funcionário", "empregado", "colaborador"]):
        return "nome"

    return None


def registro_vazio(data):
    return {
        "data": data,
        "batidas": [],
        "status_validacao": "Pendente",
        "observacoes": [],
    }


def montar_retorno(metodo, colaboradores, confianca=0, status="Processado", observacoes=None):
    return {
        "status": status,
        "metodo_extracao": metodo,
        "confianca_leitura": max(0, min(100, int(confianca))),
        "colaboradores": colaboradores,
        "observacoes": observacoes or [],
    }


def finalizar_colaborador(nome, registros, metodo, confianca):
    registros_final = []

    for registro in registros:
        batidas_validas = [b for b in registro["batidas"] if b.get("horario") and b.get("horario") != "00:00"]
        registro["status_validacao"] = "OK" if len(batidas_validas) >= 2 else "Revisar"
        if len(batidas_validas) < 2:
            registro["observacoes"].append("Menos de duas batidas identificadas no dia.")
        registros_final.append(registro)

    status_colaborador = "OK" if registros_final and all(r["status_validacao"] == "OK" for r in registros_final) else "Revisar"
    competencia = ""
    if registros_final:
        partes = registros_final[0]["data"].split("/")
        if len(partes) == 3:
            competencia = f"{partes[1]}/{partes[2]}"

    return {
        "nome_funcionario": nome or "",
        "competencia": competencia,
        "metodo_extracao": metodo,
        "confianca_leitura": max(0, min(100, int(confianca))),
        "assinatura": {
            "presente": False,
            "tipo": "ausente",
            "confianca": 0,
            "revisao_humana": False,
        },
        "registros": registros_final,
        "status_validacao": status_colaborador,
    }


def parse_linhas_logicas(linhas, metodo="Processamento Logico"):
    if not linhas:
        return []

    cabecalho = [normalizar_texto(c) for c in linhas[0]]
    mapa = {idx: detectar_tipo_coluna(coluna) for idx, coluna in enumerate(cabecalho)}
    registros_por_nome = {}

    for linha in linhas[1:]:
        if not any(normalizar_texto(c) for c in linha):
            continue

        nome = ""
        data = ""
        batidas = []

        for idx, valor in enumerate(linha):
            tipo = mapa.get(idx)
            if tipo == "nome":
                nome = normalizar_texto(valor)
            elif tipo == "data":
                data = normalizar_data(valor)
            elif tipo in TIPOS_BATIDA:
                horario = normalizar_horario(valor)
                batidas.append({
                    "tipo": tipo,
                    "horario": horario,
                    "confianca": 95 if horario else 30,
                    "revisao_humana": not bool(horario),
                })

        if not data:
            texto_linha = " ".join(normalizar_texto(c) for c in linha)
            data = normalizar_data(texto_linha)

        if not batidas:
            horarios = re.findall(r"\b(?:[01]?\d|2[0-3])\s*[:hH;.]\s*[0-5]\d\b", " ".join(map(str, linha)))
            batidas = [
                {
                "tipo": f"batida_{i + 1}",
                "horario": normalizar_horario(h),
                "confianca": 40 if normalizar_horario(h) == "00:00" else 85,
                "revisao_humana": normalizar_horario(h) == "00:00",
                }
                for i, h in enumerate(horarios)
            ]

        if data and batidas:
            nome_key = nome or "Funcionario nao identificado"
            registros_por_nome.setdefault(nome_key, []).append({
                **registro_vazio(data),
                "batidas": batidas,
            })

    return [
        finalizar_colaborador(nome, registros, metodo, 92)
        for nome, registros in registros_por_nome.items()
    ]


def ler_csv(caminho):
    with open(caminho, "r", encoding="utf-8-sig", newline="") as arquivo:
        sample = arquivo.read(2048)
        arquivo.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
        return list(csv.reader(arquivo, dialect))


def ler_txt(caminho):
    with open(caminho, "r", encoding="utf-8", errors="ignore") as arquivo:
        return arquivo.read()


def ler_excel(caminho):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Instale openpyxl para ler Excel: pip install openpyxl") from exc

    workbook = load_workbook(caminho, data_only=True, read_only=True)
    linhas = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            linhas.append([cell if cell is not None else "" for cell in row])
    return linhas


def extrair_nome_do_texto(texto):
    padroes = [
        r"(?:Nome|Funcion[áa]rio|Empregado|Colaborador)\s*[:\-]?\s*\n?\s*([A-ZÀ-Ú][A-ZÀ-Úa-zà-ú' ]{5,80})",
        r"\bNome\s*\n\s*([A-ZÀ-Ú][A-ZÀ-Úa-zà-ú' ]{5,80})",
    ]
    bloqueios = ["HORARIO", "HORÁRIO", "TRABALHO", "EMPRESA", "CNPJ", "FOLHA"]

    for padrao in padroes:
        match = re.search(padrao, texto, re.IGNORECASE)
        if match:
            nome = normalizar_texto(match.group(1)).title()
            if not any(b in nome.upper() for b in bloqueios):
                return nome

    return ""


def parse_texto_universal(texto, metodo, confianca_base=75):
    nome = extrair_nome_do_texto(texto)
    registros = []

    for match in re.finditer(r"\b([0-3]?\d[/.\-][01]?\d[/.\-]\d{2,4})\b", texto):
        data = normalizar_data(match.group(1))
        trecho = texto[match.end():match.end() + 220]
        horarios = re.findall(r"\b(?:[01]?\d|2[0-3])\s*[:hH;.]\s*[0-5]\d\b", trecho)
        batidas = []
        for i, horario_raw in enumerate(horarios[:6]):
            horario = normalizar_horario(horario_raw)
            batidas.append({
                "tipo": f"batida_{i + 1}",
                "horario": horario,
                "confianca": 40 if horario == "00:00" else (confianca_base if horario else 35),
                "revisao_humana": (not bool(horario)) or horario == "00:00",
            })

        if data and batidas:
            registros.append({
                **registro_vazio(data),
                "batidas": batidas,
            })

    if not registros:
        return []

    return [finalizar_colaborador(nome, registros, metodo, confianca_base)]


def texto_pdf_nativo(caminho):
    with fitz.open(caminho) as doc:
        return "\n".join(page.get_text() for page in doc)


def renderizar_pdf_pymupdf(caminho, dpi=220):
    imagens = []
    zoom = dpi / 72
    with fitz.open(caminho) as doc:
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            imagens.append(arr)
    return imagens


def preprocessar_para_ocr(imagem):
    arr = np.array(imagem)
    if arr.ndim == 3:
        gray = cv2.cvtColor(arr, cv2.COLOR_RGB2GRAY)
    else:
        gray = arr

    gray = cv2.resize(gray, None, fx=1.35, fy=1.35, interpolation=cv2.INTER_CUBIC)
    gray = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(gray)
    return cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 2)


def ocr_imagens(imagens):
    textos = []
    for imagem in imagens:
        proc = preprocessar_para_ocr(imagem)
        texto = pytesseract.image_to_string(
            proc,
            lang="por+eng",
            config="--oem 3 --psm 6 -c preserve_interword_spaces=1",
        )
        textos.append(texto)
    return "\n".join(textos)


def inferir_com_ia(caminho=None, texto=None):
    try:
        import google.generativeai as genai

        modelo = modelos_disponiveis()[0]
        model = genai.GenerativeModel(modelo)

        partes = [SYSTEM_PROMPT_CARTAO_PONTO]
        if caminho and Path(caminho).suffix.lower() in EXT_PDF:
            partes.append(genai.upload_file(caminho, mime_type="application/pdf"))
        else:
            partes.append(f"TEXTO OCR OU CONTEUDO ESTRUTURADO:\n{texto or ''}")

        response = model.generate_content(
            partes,
            generation_config={
                "temperature": 0.05,
                "response_mime_type": "application/json",
            },
        )
        dados = json.loads(limpar_json(response.text))
        return normalizar_resposta_ia(dados)
    except Exception as exc:
        return montar_retorno(
            "Inferencia de IA",
            [],
            0,
            status="Erro",
            observacoes=[f"Falha na inferencia de IA: {exc}"],
        )


def normalizar_resposta_ia(dados):
    colaboradores = []
    for pessoa in dados.get("colaboradores", []):
        registros = []
        for registro in pessoa.get("registros", []):
            batidas = []
            for batida in registro.get("batidas", []):
                confianca = int(batida.get("confianca", pessoa.get("confianca_leitura", 70)) or 0)
                batidas.append({
                    "tipo": batida.get("tipo", ""),
                    "horario": normalizar_horario(batida.get("horario", "")),
                    "confianca": confianca,
                    "revisao_humana": bool(batida.get("revisao_humana", confianca < 80)),
                })
            data = normalizar_data(registro.get("data", ""))
            if data:
                registros.append({
                    **registro_vazio(data),
                    "batidas": batidas,
                })
        colaboradores.append(finalizar_colaborador(
            pessoa.get("nome_funcionario") or pessoa.get("nome") or "",
            registros,
            "Inferencia de IA",
            pessoa.get("confianca_leitura", dados.get("confianca_leitura", 70)),
        ))
        assinatura = pessoa.get("assinatura")
        if isinstance(assinatura, dict):
            colaboradores[-1]["assinatura"] = {
                "presente": bool(assinatura.get("presente")),
                "tipo": assinatura.get("tipo") or ("rubrica" if assinatura.get("presente") else "ausente"),
                "confianca": max(0, min(100, int(assinatura.get("confianca", 0) or 0))),
                "revisao_humana": bool(assinatura.get("revisao_humana", False)),
            }

    return montar_retorno(
        "Inferencia de IA",
        colaboradores,
        dados.get("confianca_leitura", 70),
        status=dados.get("status", "Processado"),
        observacoes=dados.get("observacoes", []),
    )


def processar_cartao_ponto_universal(caminho):
    caminho = str(caminho)
    ext = Path(caminho).suffix.lower()

    if ext in {".csv"}:
        colaboradores = parse_linhas_logicas(ler_csv(caminho))
        return montar_retorno("Processamento Logico", colaboradores, 92)

    if ext in {".xlsx", ".xlsm"}:
        colaboradores = parse_linhas_logicas(ler_excel(caminho))
        return montar_retorno("Processamento Logico", colaboradores, 92)

    if ext == ".txt":
        texto = ler_txt(caminho)
        colaboradores = parse_texto_universal(texto, "Processamento Logico", 88)
        if colaboradores:
            return montar_retorno("Processamento Logico", colaboradores, 88)
        return inferir_com_ia(texto=texto)

    if ext in EXT_PDF:
        texto = texto_pdf_nativo(caminho)
        if len(texto.strip()) > 200:
            colaboradores = parse_texto_universal(texto, "Processamento Logico", 84)
            if colaboradores:
                return montar_retorno("Processamento Logico", colaboradores, 84)

        texto_ocr = ocr_imagens(renderizar_pdf_pymupdf(caminho))
        colaboradores = parse_texto_universal(texto_ocr, "OCR Tradicional", 72)
        if colaboradores:
            return montar_retorno("OCR Tradicional", colaboradores, 72)

        return inferir_com_ia(caminho=caminho, texto=texto_ocr)

    if ext in EXT_IMAGEM:
        imagem = Image.open(caminho)
        texto_ocr = ocr_imagens([imagem])
        colaboradores = parse_texto_universal(texto_ocr, "OCR Tradicional", 72)
        if colaboradores:
            return montar_retorno("OCR Tradicional", colaboradores, 72)
        return inferir_com_ia(texto=texto_ocr)

    return montar_retorno(
        "Processamento Logico",
        [],
        0,
        status="Erro",
        observacoes=[f"Formato nao suportado: {ext}"],
    )


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Leitor universal de cartao ponto.")
    parser.add_argument("arquivo", help="Caminho do PDF, imagem, Excel, CSV ou TXT.")
    args = parser.parse_args()

    print(json.dumps(processar_cartao_ponto_universal(args.arquivo), ensure_ascii=False, indent=2))
